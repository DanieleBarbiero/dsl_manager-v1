"""Rigenera le fixture binarie del laboratorio Orione Assistenza senza rete.

Eseguire con l'interprete Python 3.12 configurato dal progetto. Il programma
scrive soltanto dentro ``corpus_mock_orione_assistenza`` e produce ZIP OOXML
con timestamp DOS fisso 1980-01-01. Il payload VBA e' un marcatore inerte, non
codice eseguibile, inserito per verificare rilevazione e hashing.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import xlsxwriter
from docx import Document
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
ACTIVE = ROOT / "corpus" / "active"
CONTROLLED = ROOT / "materiale_di_supporto" / "fixture_controllate"
CHECKSUMS = ROOT / "materiale_di_supporto" / "checksums.json"
INVENTORY = ROOT / "materiale_di_supporto" / "inventario_fonti.csv"
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
FIXED_CREATED = datetime(2024, 7, 12, 9, 30, tzinfo=UTC)
INERT_VBA = (
    b"ORIONE_ASSISTENZA_INERT_VBA_PROJECT\x00"
    b"fixture-only; no executable VBA stream; never execute\x00"
)


ACTIVE_ROWS = (
    ("database/dump_oracle_ddl.sql", "corrente", "utile", "oracle_ddl", "estrarre", "Schema legacy con PK, FK, view e indici."),
    ("plsql/prc_assegna_tecnico.sql", "corrente", "utile", "plsql", "estrarre", "Procedura di assegnazione e dipendenze conservative."),
    ("plsql/trg_chiudi_intervento.sql", "corrente", "utile", "plsql", "estrarre_con_cautela", "Trigger supportato con regola di dominio non interamente deducibile."),
    ("forms/frm_intervento.xml", "corrente", "utile", "oracle_forms_xml", "estrarre", "Campi, pulsanti e riferimenti a tabelle."),
    ("logs/interventi_2026.log", "corrente", "utile", "log", "osservare_non_promuovere", "Eventi normali, warning, errore e timestamp con timezone nel messaggio."),
    ("documenti/nuovi_utili/requisiti_modernizzazione_2026.md", "corrente_aperto", "utile", "markdown", "interpretare_e_revisionare", "Requisiti discorsivi, asserto troppo forte e decisioni pending."),
    ("documenti/nuovi_utili/decorrenza_servizio_2026.txt", "corrente_aperto", "utile", "text", "interpretare_e_concordare", "Seconda attestazione indipendente di decorrenza e checklist."),
    ("documenti/nuovi_utili/manuale_operativo_2026.docx", "corrente_metadata_incoerente", "utile", "docx", "interpretare_rifiutando_metadata", "Manuale corrente con proprietà OOXML volutamente storiche."),
    ("documenti/nuovi_utili/matrice_stati_sla_2026.xlsx", "corrente", "utile", "xlsx", "estrarre_regioni", "Workbook ricco con formule, visibilità, named range e link esterno inventariato."),
    ("documenti/nuovi_utili/calcolo_sla_macro_2026.xlsm", "corrente", "utile", "xlsm", "estrarre_senza_eseguire_macro", "Workbook macro-enabled con vbaProject.bin sintetico e inerte."),
    ("documenti/nuovi_utili/architettura_integrazioni_2026.pptx", "corrente_metadata_incoerente", "utile", "pptx", "interpretare_e_lasciare_pending", "Architettura proposta, non decisione, con metadati incoerenti."),
    ("documenti/vecchi_utili/accordo_sla_2025.pdf", "storico_chiuso", "utile", "pdf", "interpretare_come_storico", "Accordo SLA storico, sostituito dal quadro 2026."),
    ("documenti/vecchi_utili/procedura_assegnazione_2023.html", "storico_chiuso", "utile", "html", "rifiutare_regola_obsoleta", "Procedura storica con intervallo esplicito e conflitto col processo corrente."),
    ("documenti/vecchi_non_utili/inventario_arredi_2022.txt", "storico", "non_utile", "text", "escludere_dal_dominio", "Rumore storico non pertinente."),
    ("documenti/nuovi_non_utili/istruzioni_parcheggio_2026.html", "corrente", "non_utile", "html", "escludere_dal_dominio", "Rumore recente non pertinente."),
)


def _zip_info(name: str, *, executable: bool = False) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(name, ZIP_EPOCH)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.create_system = 3
    info.external_attr = (0o100755 if executable else 0o100644) << 16
    return info


def _canonicalize_zip(path: Path, transforms: dict[str, bytes] | None = None) -> None:
    transforms = transforms or {}
    with zipfile.ZipFile(path, "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members.update(transforms)
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(members):
            archive.writestr(_zip_info(name), members[name], compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    path.write_bytes(payload.getvalue())


def _replace_zip_member(path: Path, name: str, transform) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        current = archive.read(name)
    _canonicalize_zip(path, {name: transform(current)})


def _build_docx(path: Path) -> None:
    document = Document()
    core = document.core_properties
    core.title = "Manuale operativo Orione Assistenza"
    core.subject = "Fixture didattica"
    core.author = "Laboratorio Orione"
    core.created = FIXED_CREATED.replace(tzinfo=None)
    core.modified = FIXED_CREATED.replace(tzinfo=None)
    document.add_heading("Manuale operativo Orione Assistenza", 0)
    document.add_paragraph("Edizione applicabile dal 1 marzo 2026.")
    document.add_heading("Presa in carico", 1)
    document.add_paragraph(
        "Il coordinatore assegna ogni intervento a un tecnico attivo. "
        "L'assegnazione deve restare tracciata; l'auto-assegnazione storica non e' piu' valida."
    )
    document.add_heading("Chiusura", 1)
    document.add_paragraph(
        "Prima della chiusura l'operatore controlla le voci obbligatorie. "
        "Una voce facoltativa incompleta non impedisce la chiusura."
    )
    document.add_heading("SLA", 1)
    document.add_paragraph(
        "Per P1 il target di risoluzione e' quattro ore; la matrice costituisce il riferimento numerico. "
        "La sospensione del conteggio durante l'attesa ricambi resta da decidere."
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    _canonicalize_zip(path)


def _build_pptx(path: Path) -> None:
    presentation = Presentation()
    core = presentation.core_properties
    core.title = "Architettura integrazioni Orione 2026"
    core.subject = "Ipotesi da revisionare"
    core.author = "Laboratorio Orione"
    core.created = datetime(2023, 11, 5, 8, 0)
    core.modified = datetime(2023, 11, 5, 8, 0)
    title = presentation.slides.add_slide(presentation.slide_layouts[0])
    title.shapes.title.text = "Orione Assistenza: integrazioni proposte"
    title.placeholders[1].text = "Bozza 2026; non costituisce decisione architetturale"
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Flussi da confermare"
    body = slide.placeholders[1].text_frame
    body.text = "Oracle Forms invia la richiesta al servizio di assegnazione."
    for line in (
        "Il servizio legge asset e tecnico, poi registra l'assegnazione.",
        "Una notifica verso un sistema ticket esterno e' solo ipotizzata.",
        "Identita' del sistema, protocollo e ownership restano pending.",
    ):
        body.add_paragraph().text = line
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    _canonicalize_zip(path)


def _new_workbook(path: Path, *, macro: bool = False) -> xlsxwriter.Workbook:
    options = {"in_memory": False}
    workbook = xlsxwriter.Workbook(path, options)
    workbook.set_properties(
        {
            "title": "Orione Assistenza - stati SLA",
            "subject": "Fixture deterministica",
            "author": "Laboratorio Orione",
            "company": "Organizzazione fittizia",
            "created": FIXED_CREATED.replace(tzinfo=None),
            "comments": "Le macro non devono essere eseguite." if macro else "Nessuna rete necessaria.",
        }
    )
    return workbook


def _inject_external_link(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        content_types = archive.read("[Content_Types].xml")
        relationships = archive.read("xl/_rels/workbook.xml.rels")
        workbook_xml = archive.read("xl/workbook.xml")
    content_types = content_types.replace(
        b"</Types>",
        b'<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>',
    )
    relationships = relationships.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdOrioneExternal" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" Target="externalLinks/externalLink1.xml"/></Relationships>',
    )
    workbook_xml = workbook_xml.replace(
        b"</workbook>",
        b'<externalReferences><externalReference xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdOrioneExternal"/></externalReferences></workbook>',
    )
    external_link = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        b'<externalBook r:id="rId1"/></externalLink>'
    )
    external_rels = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
        b'Target="https://orione.invalid/supporto_sla.xlsx" TargetMode="External"/></Relationships>'
    )
    _canonicalize_zip(
        path,
        {
            "[Content_Types].xml": content_types,
            "xl/_rels/workbook.xml.rels": relationships,
            "xl/workbook.xml": workbook_xml,
            "xl/externalLinks/externalLink1.xml": external_link,
            "xl/externalLinks/_rels/externalLink1.xml.rels": external_rels,
        },
    )


def _remove_formula_cache(path: Path, cell_ref: str) -> None:
    pattern = re.compile(
        rb'(<c\b[^>]*\br="' + re.escape(cell_ref.encode("ascii")) + rb'"[^>]*>.*?<f[^>]*>.*?</f>)<v>.*?</v>(</c>)',
        re.DOTALL,
    )

    def transform(data: bytes) -> bytes:
        changed, count = pattern.subn(rb"\1\2", data, count=1)
        if count != 1:
            raise RuntimeError(f"Formula cache for {cell_ref} was not found.")
        return changed

    _replace_zip_member(path, "xl/worksheets/sheet1.xml", transform)


def _build_xlsx(path: Path, *, partial_variant: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = _new_workbook(path)
    summary = workbook.add_worksheet("stati_sla")
    detail = workbook.add_worksheet("dettaglio_tecnico")
    internal = workbook.add_worksheet("parametri_interni")
    detail.hide()
    internal.very_hidden()

    header = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "border": 1})
    date_format = workbook.add_format({"num_format": "yyyy-mm-dd"})
    summary.write_row("A1", ["priorita", "stato", "ore_risoluzione", "attiva", "valida_dal"], header)
    rows = (
        ("P1", "ATTIVO", 4, True, datetime(2026, 3, 1)),
        ("P2", "ATTIVO", 12, True, datetime(2026, 3, 1)),
        ("P3", "MONITORATO", 36, False, datetime(2026, 3, 1)),
    )
    for row_index, row in enumerate(rows, 1):
        summary.write(row_index, 0, row[0])
        summary.write(row_index, 1, row[1])
        summary.write_number(row_index, 2, row[2])
        summary.write_boolean(row_index, 3, row[3])
        summary.write_datetime(row_index, 4, row[4], date_format)
    summary.add_table("A1:E4", {"name": "tabella_sla", "columns": [{"header": value} for value in rows and ["priorita", "stato", "ore_risoluzione", "attiva", "valida_dal"]]})
    summary.merge_range("A6:B6", "controlli formula")
    summary.write("A7", "doppio target P1")
    summary.write_formula("B7", "=C2*2", None, 8)
    summary.write("A8", "formula senza cache")
    summary.write_formula("B8", "=C3+C4", None, 48)
    summary.write("A9", "errore controllato")
    summary.write_formula("B9", "=1/0", None, "#DIV/0!")
    summary.write_blank("C9", None)
    summary.write("D9", "fine regione")
    summary.write("G1", "regione_secondaria")
    summary.write("G2", "sospensione_attesa_ricambi")
    summary.write("H2", "PENDING")
    if partial_variant:
        summary.write("G4", "fixture_controllata")
        summary.write("H4", "partial_success solo con worker iniettato")

    detail.write_row("A1", ["codice", "descrizione", "valore"], header)
    detail.write_row("A2", ["WARN_P1", "avviso prima della scadenza", 0.75])
    detail.write_row("A3", ["TZ", "timezone operativa", "Europe/Rome"])
    internal.write_row("A1", ["chiave", "valore"], header)
    internal.write_row("A2", ["versione_policy", "orione_sla_v_01"])

    workbook.define_name("target_p1_ore", "=stati_sla!$C$2")
    workbook.define_name("stati_sla!area_stati", "=stati_sla!$A$1:$E$4")
    workbook.close()
    _remove_formula_cache(path, "B8")
    _inject_external_link(path)


def _inject_inert_vba(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        content_types = archive.read("[Content_Types].xml")
        relationships = archive.read("xl/_rels/workbook.xml.rels")
    content_types = content_types.replace(
        b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
        b'application/vnd.ms-excel.sheet.macroEnabled.main+xml',
    ).replace(
        b"</Types>",
        b'<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
    )
    relationships = relationships.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdOrioneVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
    )
    _canonicalize_zip(
        path,
        {
            "[Content_Types].xml": content_types,
            "xl/_rels/workbook.xml.rels": relationships,
            "xl/vbaProject.bin": INERT_VBA,
        },
    )


def _build_xlsm(path: Path) -> None:
    temporary = path.with_suffix(".xlsx")
    _build_xlsx(temporary)
    path.write_bytes(temporary.read_bytes())
    temporary.unlink()
    _inject_inert_vba(path)


def _pdf_string(value: str) -> bytes:
    return ("<" + value.encode("cp1252", errors="replace").hex().upper() + ">").encode("ascii")


def _build_pdf(path: Path) -> None:
    lines = [
        "Accordo SLA Orione Assistenza 2025",
        "Validita di dominio: dal 2025-01-01 al 2025-12-31.",
        "P1: presa in carico entro 30 minuti e risoluzione entro 6 ore.",
        "Questo accordo e storico: i valori 2026 lo sostituiscono.",
        "L'attesa ricambi sospendeva il conteggio; la regola corrente e in discussione.",
    ]
    commands = [b"BT", b"/F1 12 Tf", b"72 760 Td"]
    for index, line in enumerate(lines):
        if index:
            commands.append(b"0 -22 Td")
        commands.append(_pdf_string(line) + b" Tj")
    commands.append(b"ET")
    stream = b"\n".join(commands) + b"\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Title " + _pdf_string("Accordo SLA Orione 2025") + b" /Author " + _pdf_string("Laboratorio Orione") + b" /CreationDate (D:20250101090000+01'00') /ModDate (D:20251231180000+01'00') >>",
    ]
    data = bytearray(b"%PDF-1.4\n%\xE2\xE3\xCF\xD3\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode("ascii"))
        data.extend(obj)
        data.extend(b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    data.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info 6 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(data))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_manifests() -> None:
    controlled_rows = (
        ("fixture_controllate/ai_response_orione_assistenza_controllata.jsonl", "replay", "controllata", "jsonl", "verificare_prima_del_replay", "Risposta AI congelata; non e' una chiamata a un modello."),
        ("fixture_controllate/workbook_malformed_controllato.xlsx", "controllato", "non_operativa", "xlsx_malformed", "attendere_rifiuto", "ZIP volutamente troncato per il percorso di rifiuto."),
        ("fixture_controllate/workbook_partial_controllato.xlsx", "controllato", "non_operativa", "xlsx_partial", "usare_solo_con_worker_controllato", "Workbook valido per il contratto partial_success iniettato, non fonte operativa."),
    )
    rows: list[dict[str, str]] = []
    checksums: dict[str, str] = {}
    for relative, period, utility, kind, action, reason in ACTIVE_ROWS:
        absolute = ACTIVE / Path(relative)
        if not absolute.is_file():
            raise FileNotFoundError(absolute)
        public_path = f"corpus/active/{relative}"
        digest = _sha256(absolute)
        checksums[public_path] = digest
        rows.append({"percorso": public_path, "periodo": period, "utilita": utility, "tipo": kind, "azione_attesa": action, "sha256": digest, "motivazione": reason})
    for relative, period, utility, kind, action, reason in controlled_rows:
        absolute = ROOT / "materiale_di_supporto" / relative
        if not absolute.is_file():
            raise FileNotFoundError(absolute)
        public_path = f"materiale_di_supporto/{relative}"
        digest = _sha256(absolute)
        checksums[public_path] = digest
        rows.append({"percorso": public_path, "periodo": period, "utilita": utility, "tipo": kind, "azione_attesa": action, "sha256": digest, "motivazione": reason})
    CHECKSUMS.write_text(json.dumps({"algorithm": "sha256", "files": dict(sorted(checksums.items()))}, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n")
    with INVENTORY.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=("percorso", "periodo", "utilita", "tipo", "azione_attesa", "sha256", "motivazione"), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    _build_docx(ACTIVE / "documenti/nuovi_utili/manuale_operativo_2026.docx")
    _build_xlsx(ACTIVE / "documenti/nuovi_utili/matrice_stati_sla_2026.xlsx")
    _build_xlsm(ACTIVE / "documenti/nuovi_utili/calcolo_sla_macro_2026.xlsm")
    _build_pptx(ACTIVE / "documenti/nuovi_utili/architettura_integrazioni_2026.pptx")
    _build_pdf(ACTIVE / "documenti/vecchi_utili/accordo_sla_2025.pdf")
    CONTROLLED.mkdir(parents=True, exist_ok=True)
    malformed = CONTROLLED / "workbook_malformed_controllato.xlsx"
    malformed.write_bytes(b"PK\x03\x04ORIONE_ASSISTENZA_MALFORMED_CONTROLLED\n")
    _build_xlsx(CONTROLLED / "workbook_partial_controllato.xlsx", partial_variant=True)
    ai_response = CONTROLLED / "ai_response_orione_assistenza_controllata.jsonl"
    if not ai_response.exists():
        ai_response.write_text("", encoding="utf-8", newline="\n")
    _write_manifests()
    print(json.dumps({"status": "completed", "root": str(ROOT), "checksums": str(CHECKSUMS), "inventory": str(INVENTORY)}, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
