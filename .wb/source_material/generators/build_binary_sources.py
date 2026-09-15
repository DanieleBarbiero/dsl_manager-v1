from __future__ import annotations

from datetime import datetime
from pathlib import Path

import xlsxwriter
from docx import Document


BASE = Path(__file__).resolve().parents[1]
SPREADSHEETS = BASE / "spreadsheets"
DOCUMENTS = BASE / "documenti"


def build_xlsx() -> None:
    path = SPREADSHEETS / "pianificazione_interventi_2026.xlsx"
    workbook = xlsxwriter.Workbook(path)
    workbook.set_properties(
        {
            "title": "Pianificazione interventi Orione",
            "author": "Laboratorio manuale",
            "created": datetime(2026, 1, 15, 9, 0, 0),
            "comments": "Workbook originale del laboratorio; macro assenti.",
        }
    )
    main = workbook.add_worksheet("Piano")
    archive = workbook.add_worksheet("Archivio")
    archive.hide()
    headers = ["Intervento", "Priorità", "Ore stimate", "Ore consuntive", "Scostamento"]
    for column, value in enumerate(headers):
        main.write(0, column, value)
    rows = [
        ["INT-104", "ALTA", 4, 5],
        ["INT-105", "MEDIA", 2, 2],
        ["INT-106", "BASSA", 1, 0],
    ]
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            main.write(row_index, column, value)
        cached = row[3] - row[2]
        main.write_formula(row_index, 4, f"=D{row_index + 1}-C{row_index + 1}", None, cached)
    main.add_table("A1:E4", {"name": "InterventiPianificati", "columns": [{"header": h} for h in headers]})
    main.merge_range("G1:H1", "Soglie operative")
    main.write("G2", "Livello minimo alta priorità")
    main.write_number("H2", 3)
    workbook.define_name("SogliaLivello", "='Piano'!$H$2")
    archive.write_row("A1", ["Codice", "Descrizione"])
    archive.write_row("A2", ["ARCH-1", "Riga storica"])
    workbook.close()


def build_xlsm() -> None:
    vba_path = SPREADSHEETS / "vba_project_inerte.bin"
    vba_path.write_bytes(b"ORIONE-VBA-INERTE-NON-ESEGUIRE\x00\x01")
    path = SPREADSHEETS / "controlli_inerti_2026.xlsm"
    workbook = xlsxwriter.Workbook(path)
    workbook.set_properties(
        {
            "title": "Controlli inerti Orione",
            "author": "Laboratorio manuale",
            "created": datetime(2026, 1, 15, 9, 0, 0),
        }
    )
    sheet = workbook.add_worksheet("Controlli")
    sheet.write_row("A1", ["Codice", "Attivo"])
    sheet.write_row("A2", ["SICUREZZA", True])
    sheet.write_row("A3", ["COLLAUDO", True])
    workbook.add_vba_project(vba_path)
    workbook.close()
    vba_path.unlink()


def build_docx() -> None:
    path = DOCUMENTS / "nota_sla_orione_2026.docx"
    document = Document()
    document.core_properties.title = "Nota SLA Orione"
    document.core_properties.author = "Laboratorio manuale"
    document.core_properties.created = datetime(2026, 1, 15, 9, 0, 0)
    document.add_heading("Nota SLA Orione", level=0)
    document.add_paragraph("effective_date: 2026-01-15")
    document.add_heading("Priorità alta", level=1)
    document.add_paragraph(
        "Gli interventi ad alta priorità devono essere presi in carico entro due ore "
        "da un tecnico con livello almeno 3. La presa in carico e la chiusura sono eventi distinti."
    )
    document.add_heading("Chiusura", level=1)
    document.add_paragraph(
        "La checklist SICUREZZA e la checklist COLLAUDO devono entrambe avere esito OK. "
        "Questa nota conferma il manuale 2026 ma non chiarisce le certificazioni temporanee."
    )
    document.save(path)


if __name__ == "__main__":
    SPREADSHEETS.mkdir(parents=True, exist_ok=True)
    DOCUMENTS.mkdir(parents=True, exist_ok=True)
    build_xlsx()
    build_xlsm()
    build_docx()

